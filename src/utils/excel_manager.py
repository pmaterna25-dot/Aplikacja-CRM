"""
Excel management utilities for the Materna CRM application.
"""

import os
from typing import List, Dict, Any, Optional
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from config import EXCEL_FILE, columns
from models.client import Client, Statistics


class ExcelManager:
    """Manages Excel file operations for the CRM system."""

    def __init__(self, filename: str = EXCEL_FILE):
        self.filename = filename
        self._ensure_file_exists()

    def _ensure_file_exists(self):
        """Ensure the Excel file exists with proper structure."""
        if not os.path.exists(self.filename):
            wb = Workbook()
            ws = wb.active
            ws.title = "Klienci"
            ws.append(columns)
            wb.save(self.filename)

    def _get_main_worksheet(self) -> Worksheet:
        """Get the main clients worksheet."""
        wb = load_workbook(self.filename)
        return wb.active

    def _get_statistics_worksheet(self) -> Worksheet:
        """Get or create the statistics worksheet."""
        wb = load_workbook(self.filename)
        if "Statystyka" not in wb.sheetnames:
            ws_stat = wb.create_sheet("Statystyka")
            ws_stat.append([
                "Data", "Nowych klientów", "Odebrany", "Nieodebrany",
                "Umówione spotkanie", "Niezainteresowany", "Ponowny kontakt", "Potencjał"
            ])
            wb.save(self.filename)
        return wb["Statystyka"]

    def save_client(self, client: Client) -> bool:
        """Save a client to Excel."""
        try:
            wb = load_workbook(self.filename)
            ws = wb.active

            # Find next row number
            next_row = ws.max_row + 1

            # Prepare data
            data = [
                next_row,  # #
                client.imie_nazwisko,
                client.telefon,
                client.email,
                client.rocznica_majatku,
                client.polisa_zyciowa,
                client.notatki,
                client.nieruchomosc,
                client.status_klienta,
                client.styl_rozmowy,
                client.data_kontaktu,
                client.ponowny_kontakt,
                client.godzina_kontaktu,
                client.miejscowosc,
                client.czy_posiada_dzieci,
                client.status_rozmowy
            ]

            ws.append(data)
            wb.save(self.filename)
            return True
        except Exception as e:
            print(f"Error saving client: {e}")
            return False

    def update_client(self, row_number: int, client: Client) -> bool:
        """Update a client in Excel."""
        try:
            wb = load_workbook(self.filename)
            ws = wb.active

            excel_row = row_number + 1
            if excel_row > ws.max_row:
                return False

            # Update mapping
            mapping = {
                2: client.imie_nazwisko,
                3: client.telefon,
                4: client.email,
                5: client.rocznica_majatku,
                6: client.polisa_zyciowa,
                7: client.notatki,
                8: client.nieruchomosc,
                9: client.status_klienta,
                10: client.styl_rozmowy,
                11: client.data_kontaktu,
                12: client.ponowny_kontakt,
                13: client.godzina_kontaktu,
                14: client.miejscowosc,
                15: client.czy_posiada_dzieci,
                16: client.status_rozmowy
            }

            for col_idx, value in mapping.items():
                ws.cell(row=excel_row, column=col_idx, value=value)

            wb.save(self.filename)
            return True
        except Exception as e:
            print(f"Error updating client: {e}")
            return False

    def delete_client(self, row_number: int) -> bool:
        """Delete a client from Excel."""
        try:
            wb = load_workbook(self.filename)
            ws = wb.active

            rows = list(ws.iter_rows(values_only=True))
            header = rows[0]

            # Create new rows without the deleted one
            new_rows = [header]
            for r in rows[1:]:
                if r[0] != row_number:
                    new_rows.append(r)

            # Clear and rewrite worksheet
            ws.delete_rows(1, ws.max_row)
            for row in new_rows:
                ws.append(row)

            wb.save(self.filename)
            return True
        except Exception as e:
            print(f"Error deleting client: {e}")
            return False

    def load_all_clients(self) -> List[Client]:
        """Load all clients from Excel."""
        try:
            wb = load_workbook(self.filename)
            ws = wb.active

            clients = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:  # Skip empty rows
                    client_data = dict(zip(columns, row))
                    # Convert to Client model format
                    client_dict = {
                        'imie_nazwisko': client_data.get("Imię i nazwisko", ""),
                        'telefon': client_data.get("Telefon", ""),
                        'email': client_data.get("Email", ""),
                        'rocznica_majatku': client_data.get("Rocznica majątku", ""),
                        'polisa_zyciowa': client_data.get("Polisa życiowa", ""),
                        'notatki': client_data.get("Notatki", ""),
                        'nieruchomosc': client_data.get("Nieruchomość", ""),
                        'status_klienta': client_data.get("Status klienta", ""),
                        'styl_rozmowy': client_data.get("Styl rozmowy", ""),
                        'data_kontaktu': client_data.get("Data kontaktu", ""),
                        'ponowny_kontakt': client_data.get("Ponowny kontakt", ""),
                        'godzina_kontaktu': client_data.get("Godzina kontaktu", ""),
                        'miejscowosc': client_data.get("Miejscowość", ""),
                        'czy_posiada_dzieci': client_data.get("Czy posiada dzieci", ""),
                        'status_rozmowy': client_data.get("Status rozmowy", "")
                    }
                    clients.append(Client.from_dict(client_dict))

            return clients
        except Exception as e:
            print(f"Error loading clients: {e}")
            return []

    def get_today_clients_count(self, today_str: str) -> int:
        """Get count of clients saved today."""
        try:
            wb = load_workbook(self.filename)
            ws = wb.active

            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if str(row[10]) == today_str:  # Data kontaktu column
                    count += 1
            return count
        except Exception:
            return 0

    def add_statistics_for_today(self, today_str: str, status_rozmowy: str):
        """Add statistics entry for today."""
        try:
            wb = load_workbook(self.filename)
            ws_stat = self._get_statistics_worksheet()

            statusy = ["Odebrany", "Nieodebrany", "Umówione spotkanie", "Niezainteresowany", "Ponowny kontakt", "Potencjał"]

            # Find existing row for today
            found_row = None
            for row_idx, row in enumerate(ws_stat.iter_rows(min_row=2, values_only=False), 2):
                if str(row[0].value) == today_str:
                    found_row = row
                    break

            if found_row:
                # Update existing row
                found_row[1].value = int(found_row[1].value or 0) + 1
                if status_rozmowy in statusy:
                    idx = statusy.index(status_rozmowy)
                    found_row[2+idx].value = int(found_row[2+idx].value or 0) + 1
            else:
                # Create new row
                row_data = [today_str, 1, 0, 0, 0, 0, 0, 0]
                if status_rozmowy in statusy:
                    idx = statusy.index(status_rozmowy)
                    row_data[2+idx] = 1
                ws_stat.append(row_data)

            wb.save(self.filename)
        except Exception as e:
            print(f"Error adding statistics: {e}")

    def get_statistics_for_date(self, date_str: str) -> Optional[Statistics]:
        """Get statistics for a specific date."""
        try:
            ws_stat = self._get_statistics_worksheet()

            for row in ws_stat.iter_rows(min_row=2, values_only=True):
                if str(row[0]) == date_str:
                    return Statistics.from_list(row)
            return None
        except Exception as e:
            print(f"Error getting statistics: {e}")
            return None
